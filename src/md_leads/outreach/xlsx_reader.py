from __future__ import annotations

from pathlib import Path
from typing import Optional

from openpyxl import load_workbook

from md_leads.exporter import LEAD_COLUMNS
from md_leads.models import LeadRow

LEADS_SHEET = "Leads"


def _opt_str(value) -> Optional[str]:
    if value is None:
        return None
    s = str(value).strip()
    return s if s else None


def _opt_float(value) -> Optional[float]:
    if value is None or value == "":
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _int_or_zero(value) -> int:
    if value is None or value == "":
        return 0
    try:
        return int(value)
    except (TypeError, ValueError):
        try:
            return int(float(value))
        except (TypeError, ValueError):
            return 0


def read_leads(path: Path, top_n: int) -> list[LeadRow]:
    """Read the 'Leads' sheet, return up to top_n LeadRows sorted by lead_score desc."""
    path = Path(path)
    wb = load_workbook(path, read_only=True, data_only=True)
    if LEADS_SHEET not in wb.sheetnames:
        raise ValueError(
            f"Sheet '{LEADS_SHEET}' lipsește în {path}. "
            f"Sheet-uri găsite: {wb.sheetnames}"
        )
    ws = wb[LEADS_SHEET]
    rows_iter = ws.iter_rows(values_only=True)
    header = next(rows_iter, None)
    if header is None:
        return []
    col = {h: i for i, h in enumerate(header)}
    missing = [c for c in LEAD_COLUMNS if c not in col]
    if missing:
        raise ValueError(
            f"Coloane lipsă în {path}: {missing}. "
            f"Așteptăm exact {LEAD_COLUMNS}"
        )

    leads: list[LeadRow] = []
    for raw in rows_iter:
        if raw is None or all(v in (None, "") for v in raw):
            continue
        name = _opt_str(raw[col["nume"]]) or ""
        if not name:
            continue
        leads.append(LeadRow(
            name=name,
            category=_opt_str(raw[col["categorie"]]) or "",
            phone=_opt_str(raw[col["telefon"]]),
            address=_opt_str(raw[col["adresa"]]) or "",
            website=_opt_str(raw[col["website"]]),
            status=_opt_str(raw[col["status_website"]]) or "",
            reason=_opt_str(raw[col["motiv"]]) or "",
            lead_score=_int_or_zero(raw[col["lead_score"]]),
            google_rating=_opt_float(raw[col["google_rating"]]),
            reviews_count=_int_or_zero(raw[col["reviews_count"]]),
            google_maps_url=_opt_str(raw[col["google_maps_url"]]) or "",
        ))

    leads.sort(key=lambda l: l.lead_score, reverse=True)
    return leads[:top_n]
