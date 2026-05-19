from __future__ import annotations

from pathlib import Path
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from md_leads.models import Lead

LEAD_COLUMNS = [
    "nume", "categorie", "telefon", "adresa", "website",
    "status_website", "motiv", "lead_score",
    "google_rating", "reviews_count",
    "pagespeed_mobile", "mobile_friendly", "an_inregistrare",
    "google_maps_url", "latitudine", "longitudine",
]


def _row(lead: Lead) -> list:
    biz = lead.business
    eb = lead.enriched
    ps = eb.pagespeed
    idno = eb.idno
    return [
        biz.name,
        biz.category,
        biz.phone or "",
        biz.address,
        biz.website or "",
        lead.status.value,
        lead.reason,
        lead.lead_score,
        biz.google_rating if biz.google_rating is not None else "",
        biz.reviews_count,
        ps.performance_mobile if ps else "",
        ps.mobile_friendly if ps else "",
        idno.registration_year if idno and idno.registration_year else "",
        biz.google_maps_url,
        biz.latitude,
        biz.longitude,
    ]


def _write_sheet(ws, rows: Iterable[Lead]) -> None:
    ws.append(LEAD_COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for lead in rows:
        ws.append(_row(lead))
    if ws.max_row >= 2:
        ws.auto_filter.ref = (
            f"A1:{get_column_letter(len(LEAD_COLUMNS))}{ws.max_row}"
        )
    # Reasonable column widths
    widths = [28, 18, 18, 36, 32, 14, 36, 10, 8, 8, 10, 8, 8, 36, 10, 10]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def write_xlsx(
    leads: list[Lead],
    skipped: list[Lead],
    out_path: Path,
) -> Path:
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    sorted_leads = sorted(leads, key=lambda l: l.lead_score, reverse=True)
    sorted_skipped = sorted(skipped, key=lambda l: l.business.name)

    wb = Workbook()
    leads_ws = wb.active
    leads_ws.title = "Leads"
    skipped_ws = wb.create_sheet("Skipped (good)")
    _write_sheet(leads_ws, sorted_leads)
    _write_sheet(skipped_ws, sorted_skipped)

    wb.save(out_path)
    return out_path
