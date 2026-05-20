from pathlib import Path

import pytest
from openpyxl import Workbook

from md_leads.outreach.xlsx_reader import read_leads


def _make_xlsx(tmp_path: Path, rows: list[list]) -> Path:
    """Create a Leads-sheet XLSX with the standard header + given rows."""
    from md_leads.exporter import LEAD_COLUMNS
    wb = Workbook()
    ws = wb.active
    ws.title = "Leads"
    ws.append(LEAD_COLUMNS)
    for r in rows:
        ws.append(r)
    wb.create_sheet("Skipped (good)").append(LEAD_COLUMNS)
    path = tmp_path / "leads.xlsx"
    wb.save(path)
    return path


def test_read_leads_returns_lead_rows(tmp_path):
    path = _make_xlsx(tmp_path, [
        ["Salon X", "Salon", "+373 22 000 000", "addr", "",
         "missing", "fără website", 11,
         4.7, 50, "", "", "", "https://maps.google.com/?cid=1", 47.0, 28.8],
    ])
    leads = read_leads(path, top_n=10)
    assert len(leads) == 1
    assert leads[0].name == "Salon X"
    assert leads[0].phone == "+373 22 000 000"
    assert leads[0].website is None
    assert leads[0].status == "missing"
    assert leads[0].lead_score == 11
    assert leads[0].google_rating == 4.7
    assert leads[0].reviews_count == 50


def test_read_leads_respects_top_n(tmp_path):
    rows = [
        ["Z", "c", "+1", "a", "", "missing", "r", 5,
         4.0, 10, "", "", "", "u", 0.0, 0.0],
        ["A", "c", "+2", "a", "", "missing", "r", 10,
         4.5, 20, "", "", "", "u", 0.0, 0.0],
        ["M", "c", "+3", "a", "", "missing", "r", 8,
         4.2, 15, "", "", "", "u", 0.0, 0.0],
    ]
    path = _make_xlsx(tmp_path, rows)
    leads = read_leads(path, top_n=2)
    assert len(leads) == 2
    assert [l.name for l in leads] == ["A", "M"]


def test_read_leads_skips_skipped_sheet(tmp_path):
    """Skipped (good) sheet must be ignored even if it has rows."""
    from md_leads.exporter import LEAD_COLUMNS
    wb = Workbook()
    leads_ws = wb.active
    leads_ws.title = "Leads"
    leads_ws.append(LEAD_COLUMNS)
    leads_ws.append(["OnlyLead", "c", "+1", "a", "", "missing", "r", 10,
                     4.0, 10, "", "", "", "u", 0.0, 0.0])
    skipped_ws = wb.create_sheet("Skipped (good)")
    skipped_ws.append(LEAD_COLUMNS)
    skipped_ws.append(["GoodOne", "c", "+2", "a", "https://x.md", "good",
                       "OK", 0, 4.5, 5, "", "", "", "u", 0.0, 0.0])
    path = tmp_path / "x.xlsx"
    wb.save(path)
    leads = read_leads(path, top_n=10)
    assert [l.name for l in leads] == ["OnlyLead"]


def test_read_leads_missing_leads_sheet_raises(tmp_path):
    wb = Workbook()
    wb.active.title = "Wrong"
    path = tmp_path / "bad.xlsx"
    wb.save(path)
    with pytest.raises(ValueError, match="Leads"):
        read_leads(path, top_n=5)


def test_read_leads_zero_rows(tmp_path):
    path = _make_xlsx(tmp_path, [])
    leads = read_leads(path, top_n=10)
    assert leads == []


def test_read_leads_handles_blank_optional_fields(tmp_path):
    """Empty website/phone cells become None, not empty strings."""
    path = _make_xlsx(tmp_path, [
        ["NoPhone", "c", "", "a", "", "missing", "r", 10,
         "", "", "", "", "", "u", 0.0, 0.0],
    ])
    leads = read_leads(path, top_n=5)
    assert len(leads) == 1
    assert leads[0].phone is None
    assert leads[0].website is None
    assert leads[0].google_rating is None
    assert leads[0].reviews_count == 0
