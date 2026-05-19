from pathlib import Path

from openpyxl import load_workbook

from md_leads.exporter import write_xlsx, LEAD_COLUMNS
from md_leads.models import (
    Business, EnrichedBusiness, IdnoRecord, Lead,
    PageSpeedResult, WebsiteCheckResult, WebsiteStatus,
)


def _lead(name="X", website=None, status=WebsiteStatus.MISSING, score=10):
    biz = Business(
        place_id="x", name=name, category="cafe", phone="+373 22 000 000",
        address="addr", website=website, google_rating=4.5,
        reviews_count=22, latitude=47.0, longitude=28.8,
        google_maps_url="https://maps.google.com/?cid=x",
    )
    eb = EnrichedBusiness(
        business=biz,
        website_check=WebsiteCheckResult(
            final_url=website, status_code=200, https=True,
            is_parking=False, error=None,
        ) if website else None,
        pagespeed=PageSpeedResult(performance_mobile=72,
                                  mobile_friendly=True) if website else None,
        idno=IdnoRecord(idno="123", registration_year=2024),
    )
    return Lead(business=biz, enriched=eb, status=status,
                reason="fără website", lead_score=score)


def test_write_xlsx_creates_file_with_two_sheets(tmp_path):
    out = tmp_path / "leads.xlsx"
    leads = [_lead("A", score=10), _lead("B", score=5)]
    skipped = [_lead("C", website="https://good.md",
                     status=WebsiteStatus.GOOD, score=0)]

    path = write_xlsx(leads, skipped, out)
    assert path.exists()

    wb = load_workbook(path)
    assert wb.sheetnames == ["Leads", "Skipped (good)"]


def test_write_xlsx_columns_match_spec(tmp_path):
    out = tmp_path / "leads.xlsx"
    write_xlsx([_lead("A")], [], out)

    wb = load_workbook(out)
    ws = wb["Leads"]
    header = [cell.value for cell in ws[1]]
    assert header == LEAD_COLUMNS


def test_write_xlsx_sorts_by_lead_score_desc(tmp_path):
    out = tmp_path / "leads.xlsx"
    leads = [_lead("low", score=3), _lead("high", score=10),
             _lead("mid", score=7)]
    write_xlsx(leads, [], out)

    wb = load_workbook(out)
    ws = wb["Leads"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    names = [r[0] for r in rows]
    assert names == ["high", "mid", "low"]


def test_write_xlsx_handles_no_skipped(tmp_path):
    out = tmp_path / "leads.xlsx"
    write_xlsx([_lead("A")], [], out)
    wb = load_workbook(out)
    assert wb["Skipped (good)"].max_row >= 1  # at least header
